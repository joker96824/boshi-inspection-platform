"""
导航控制器配置模板生成脚本
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import os

def create_navigation_controller_template():
    """创建导航控制器配置模板"""
    
    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "导航控制器配置模板"
    
    # 设置样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # 设置边框
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 定义列
    columns = [
        {"name": "模块组", "description": "模块组编号(1-10)"},
        {"name": "以太网IP", "description": "以太网通讯IP地址，如：192.168.1.100"},
        {"name": "子网掩码", "description": "子网掩码，如：255.255.255.0"},
        {"name": "网关地址", "description": "网关地址，如：192.168.1.1"},
        {"name": "以太网端口", "description": "以太网通讯端口号(1-65535)"},
        {"name": "波特率", "description": "波特率，可选值：9600,19200,38400,57600,115200,230400,460800,921600"},
        {"name": "减速距离", "description": "减速距离(mm)，范围：0-10000"},
        {"name": "停止距离", "description": "停止距离(mm)，范围：0-10000，必须大于减速距离"},
        {"name": "最大线速度", "description": "路径最大线速度(m/s)，范围：0.001-10.000"},
        {"name": "最大角速度", "description": "路径最大角速度(rad/s)，范围：0.001-10.000"},
        {"name": "加速度", "description": "加速度(m/s²)，范围：0.001-10.000"},
        {"name": "减速度", "description": "减速度(m/s²)，范围：0.001-10.000"},
        {"name": "膨胀系数", "description": "膨胀系数，范围：0-2.0"}
    ]
    
    # 写入表头
    for col_idx, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=column["name"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        
        # 设置列宽
        ws.column_dimensions[cell.column_letter].width = 20
    
    # 添加说明行（在第二行）
    ws.cell(row=2, column=1, value="说明：")
    ws.cell(row=2, column=1).font = Font(bold=True)
    
    # 添加列说明（在第三行）
    for col_idx, column in enumerate(columns, 1):
        if col_idx == 1:
            ws.cell(row=3, column=col_idx, value=f"{column['name']}: {column['description']}")
        else:
            ws.cell(row=3, column=col_idx, value=column['description'])
    
    # 设置说明行样式
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = Font(italic=True, color="666666")
        cell.border = thin_border
    
    # 添加示例数据（从第4行开始）
    example_data = [
        [1, "192.168.1.100", "255.255.255.0", "192.168.1.1", 8080, 9600, 800.0, 1500.0, 0.8, 0.1, 0.8, 0.8, 0.0],
        [2, "192.168.1.101", "255.255.255.0", "192.168.1.1", 8081, 115200, 1000.0, 2000.0, 1.0, 0.2, 1.0, 1.0, 0.1]
    ]
    
    # 写入示例数据（从第4行开始）
    for row_idx, row_data in enumerate(example_data, 4):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
    
    # 添加示例说明（在第6行）
    ws.cell(row=6, column=1, value="示例数据（可删除）：")
    ws.cell(row=6, column=1).font = Font(bold=True, color="FF0000")
    
    # 保存文件
    template_path = "templates/导航控制器配置模板.xlsx"
    wb.save(template_path)
    print(f"模板文件已创建: {template_path}")

if __name__ == "__main__":
    create_navigation_controller_template()
