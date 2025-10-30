"""
Excel处理工具类
"""

import io
import pandas as pd
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from fastapi.responses import StreamingResponse
from urllib.parse import quote
from ..config.logging import get_logger

logger = get_logger(__name__)


class ExcelProcessor:
    """Excel处理器"""
    
    @staticmethod
    def create_template(columns: List[Dict[str, str]], sheet_name: str = "模板") -> StreamingResponse:
        """创建Excel模板
        
        Args:
            columns: 列定义列表，每个字典包含name和description
            sheet_name: 工作表名称
            
        Returns:
            StreamingResponse: Excel文件流
        """
        try:
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            # 设置样式
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            # 设置边框
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # 写入表头
            for col_idx, column in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=column["name"])
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
                
                # 设置列宽
                ws.column_dimensions[cell.column_letter].width = 20
            
            # 添加说明行
            ws.cell(row=2, column=1, value="说明：")
            ws.cell(row=2, column=1).font = Font(bold=True)
            
            # 添加列说明
            for col_idx, column in enumerate(columns, 1):
                if col_idx == 1:
                    ws.cell(row=3, column=col_idx, value=f"{column['name']}: {column.get('description', '')}")
                else:
                    ws.cell(row=3, column=col_idx, value=column.get('description', ''))
            
            # 设置说明行样式
            for col_idx in range(1, len(columns) + 1):
                cell = ws.cell(row=3, column=col_idx)
                cell.font = Font(italic=True, color="666666")
                cell.border = thin_border
            
            # 创建内存流
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            return StreamingResponse(
                io.BytesIO(output.read()),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(sheet_name)}.xlsx"}
            )
            
        except Exception as e:
            logger.error(f"创建Excel模板失败: {str(e)}")
            raise
    
    @staticmethod
    def export_data(data: List[Dict[str, Any]], columns: List[Dict[str, str]], 
                   sheet_name: str = "数据导出") -> StreamingResponse:
        """导出数据到Excel（包含说明行，可直接重新导入）
        
        Args:
            data: 要导出的数据列表
            columns: 列定义列表
            sheet_name: 工作表名称
            
        Returns:
            StreamingResponse: Excel文件流
        """
        try:
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            # 设置样式
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            # 设置边框
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # 写入表头（第1行）
            for col_idx, column in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=column["name"])
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
                
                # 设置列宽
                ws.column_dimensions[cell.column_letter].width = 20
            
            # 添加说明行（第2行）
            ws.cell(row=2, column=1, value="说明：")
            ws.cell(row=2, column=1).font = Font(bold=True)
            
            # 添加列说明（第3行）
            for col_idx, column in enumerate(columns, 1):
                if col_idx == 1:
                    ws.cell(row=3, column=col_idx, value=f"{column['name']}: {column.get('description', '')}")
                else:
                    ws.cell(row=3, column=col_idx, value=column.get('description', ''))
            
            # 设置说明行样式
            for col_idx in range(1, len(columns) + 1):
                cell = ws.cell(row=3, column=col_idx)
                cell.font = Font(italic=True, color="666666")
                cell.border = thin_border
            
            # 写入数据（从第4行开始）
            if data:
                # 重命名列
                column_mapping = {col["key"]: col["name"] for col in columns if "key" in col}
                
                for row_idx, row_data in enumerate(data, 4):
                    for col_idx, column in enumerate(columns, 1):
                        if "key" in column:
                            value = row_data.get(column["key"])
                            cell = ws.cell(row=row_idx, column=col_idx, value=value)
                            cell.border = thin_border
            
            # 创建内存流
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            return StreamingResponse(
                io.BytesIO(output.read()),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(sheet_name)}.xlsx"}
            )
            
        except Exception as e:
            logger.error(f"导出Excel数据失败: {str(e)}")
            raise
    
    @staticmethod
    def parse_excel(file_content: bytes, sheet_name: str = None) -> List[Dict[str, Any]]:
        """解析Excel文件
        
        Args:
            file_content: Excel文件内容
            sheet_name: 工作表名称，None表示使用第一个工作表
            
        Returns:
            List[Dict[str, Any]]: 解析后的数据列表
        """
        try:
            logger.info(f"开始解析Excel文件，文件大小: {len(file_content)} bytes")
            
            # 读取Excel文件
            result = pd.read_excel(io.BytesIO(file_content), sheet_name=sheet_name)
            
            logger.info(f"Excel文件读取成功，数据类型: {type(result)}")
            
            # 处理不同的返回类型
            if isinstance(result, dict):
                # 多个工作表的情况，使用第一个工作表
                logger.info(f"检测到多个工作表: {list(result.keys())}")
                if sheet_name is None:
                    # 使用第一个工作表
                    sheet_name = list(result.keys())[0]
                    logger.info(f"使用第一个工作表: {sheet_name}")
                df = result[sheet_name]
            elif isinstance(result, pd.DataFrame):
                df = result
            else:
                logger.error(f"读取Excel文件失败，返回类型: {type(result)}")
                raise ValueError("Excel文件格式不正确")
            
            logger.info(f"DataFrame形状: {df.shape}, 列名: {list(df.columns)}")
            
            # 转换为字典列表
            data = df.to_dict('records')
            
            logger.info(f"转换为字典列表成功，共{len(data)}行数据")
            
            # 清理空值
            cleaned_data = []
            for idx, row in enumerate(data):
                cleaned_row = {k: v for k, v in row.items() if pd.notna(v)}
                
                # 跳过说明行和提示行
                if cleaned_row:
                    # 检查是否是说明行（包含"说明"、"示例"等关键词）
                    first_value = str(list(cleaned_row.values())[0]).strip()
                    if any(keyword in first_value for keyword in ["说明", "示例", "可删除", "模板"]):
                        logger.debug(f"跳过说明行: {idx + 1} - {first_value}")
                        continue
                    
                    # 检查是否所有值都是字符串且包含说明性文字
                    all_strings = all(isinstance(v, str) for v in cleaned_row.values())
                    if all_strings and any(keyword in str(v) for v in cleaned_row.values() for keyword in ["范围", "可选值", "如：", "必须"]):
                        logger.debug(f"跳过说明行: {idx + 1} - 包含说明性文字")
                        continue
                    
                    cleaned_data.append(cleaned_row)
                else:
                    logger.debug(f"跳过空行: {idx + 1}")
            
            logger.info(f"数据清理完成，有效数据行数: {len(cleaned_data)}")
            return cleaned_data
            
        except Exception as e:
            logger.error(f"解析Excel文件失败: {str(e)}", exc_info=True)
            raise
    
    @staticmethod
    def validate_data(data: List[Dict[str, Any]], required_fields: List[str]) -> List[str]:
        """验证数据
        
        Args:
            data: 要验证的数据
            required_fields: 必填字段列表
            
        Returns:
            List[str]: 错误信息列表
        """
        errors = []
        
        for idx, row in enumerate(data, 1):
            row_errors = []
            
            # 检查必填字段
            for field in required_fields:
                if field not in row or pd.isna(row.get(field)) or str(row.get(field)).strip() == "":
                    row_errors.append(f"第{idx}行缺少必填字段: {field}")
            
            if row_errors:
                errors.extend(row_errors)
        
        return errors
